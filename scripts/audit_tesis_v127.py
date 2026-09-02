#!/usr/bin/env python3
"""Auditoría v127 vs Devolución 4, pendientes v126, diff v126→v127."""
import re
import json
from pathlib import Path
from difflib import unified_diff

from docx import Document
from pypdf import PdfReader

PDF_PATH = Path(r"d:\Devolucion_4_Nodo_Propiedades.pdf")
V127_PATH = Path(r"d:\Nodo_Propiedades_Tesis_v127.docx")
V126_PATH = Path(r"d:\Nodo_Propiedades_Tesis_v126.docx")
OUT_DIR = Path(r"D:\Dev\Nodo-Propiedades\scripts\audit_output")
OUT_DIR.mkdir(exist_ok=True)


def extract_docx_text(path: Path) -> str:
    doc = Document(str(path))
    parts = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_docx_paragraphs(path: Path):
    doc = Document(str(path))
    paragraphs = []
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if text:
            style = p.style.name if p.style else ""
            paragraphs.append({"i": i, "style": style, "text": text})
    return paragraphs


def extract_pdf_text(path: Path) -> str:
    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        pages.append(f"--- PAGE {i+1} ---\n{text}")
    return "\n".join(pages)


def find_context(paragraphs, needle, context=120):
    results = []
    for p in paragraphs:
        if needle.lower() in p["text"].lower():
            t = p["text"]
            idx = t.lower().find(needle.lower())
            start = max(0, idx - 60)
            end = min(len(t), idx + len(needle) + 60)
            results.append(t[start:end])
    return results


def section_paragraphs(paragraphs, marker_patterns):
    out = []
    in_section = False
    for p in paragraphs:
        t = p["text"]
        if any(re.search(pat, t, re.I) for pat in marker_patterns):
            in_section = True
            out.append(t)
            continue
        if in_section:
            if re.match(r"^\d+\.\s", t) and not any(re.search(pat, t, re.I) for pat in marker_patterns):
                if re.match(r"^\d+\.\d", t) is None or len(out) > 3:
                    break
            out.append(t)
    return "\n".join(out)


def context(text: str, pos: int, radius=120) -> str:
    start = max(0, pos - radius)
    end = min(len(text), pos + radius)
    return text[start:end].replace("\n", " ")


def find_all(pattern: str, text: str, flags=0):
    return list(re.finditer(pattern, text, flags))


def run_devolucion4_checks(paragraphs, full_text):
    full_lower = full_text.lower()
    checks = {}

    # === PLAN 10 ITEMS ===
    checks["plan1_2.1_notificacion"] = {
        "desc": "Plan 1: §2.1 notificación activa al operador",
        "ok": "notificación activa" in full_lower and "mecanismo explícito" in full_lower,
        "evidence": find_context(paragraphs, "notificación activa")[:2],
    }
    checks["plan2_5.6_n6"] = {
        "desc": "Plan 2: §5.6 n=6 unificado o declaración no instrumentada",
        "ok": bool(re.search(r"n\s*=\s*6|seis ejecuciones|no instrumentada|no quedó instrumentada", full_lower)),
        "evidence": find_context(paragraphs, "5.6")[:1] + find_context(paragraphs, "instrumentada")[:1],
    }
    checks["plan2_5.6_enumerar"] = {
        "desc": "Plan 2 (parte): enumerar 6 comparaciones o retirar cifra",
        "ok": bool(re.search(r"(enumer|las seis|6 conversaciones|CAL-|TIB-|AMB-|no instrumentada|retira)", full_lower)),
        "evidence": find_context(paragraphs, "seis")[:2],
    }
    checks["plan3_tabla7_fechas"] = {
        "desc": "Plan 3: nota Tabla 7 fechas MCP/repo",
        "ok": "5 de agosto" in full_lower and ("21 de agosto" in full_lower or "repositorio" in full_lower),
        "evidence": find_context(paragraphs, "5 de agosto")[:1],
    }
    checks["plan3_figuras_3_6"] = {
        "desc": "Plan 3: Figuras 3 a 6 (no 5 a 8)",
        "ok": ("figuras 3 a 6" in full_lower or "figuras 3 a 6" in full_lower.replace("–", " a ")) and "figuras 5 a 8" not in full_lower,
        "evidence": find_context(paragraphs, "Figuras 3")[:1] or find_context(paragraphs, "figuras 3")[:1],
    }
    checks["plan4_fredes"] = {
        "desc": "Plan 4: cita (Fredes, 2026)",
        "ok": "(fredes, 2026)" in full_lower or "fredes, 2026" in full_lower,
        "evidence": find_context(paragraphs, "Fredes, 2026")[:2],
    }
    checks["plan5_repo"] = {
        "desc": "Plan 5: repositorio público declarado",
        "ok": "github.com" in full_lower or "repositorio" in full_lower,
        "evidence": find_context(paragraphs, "repositorio")[:2],
    }
    checks["plan6_figura1"] = {
        "desc": "Plan 6: Figura 1 rama notificación caliente",
        "ok": "temperatura caliente" in full_lower and ("telegram" in full_lower or "alerta" in full_lower),
        "evidence": find_context(paragraphs, "Temperatura Caliente")[:2],
    }
    checks["plan7_resumen_matiz"] = {
        "desc": "Plan 7: matiz Resumen (experiencia Mendoza, condicional)",
        "ok": ("experiencia del autor" in full_lower or "mendoza" in full_lower) and "concentrarían" in full_lower,
        "evidence": find_context(paragraphs, "concentrarían")[:1],
    }
    checks["plan7_keywords"] = {
        "desc": "Plan 7 / Q6: palabras clave unificadas",
        "ok": full_lower.count("panel comercial") >= 1,
        "evidence": [f"'panel comercial' aparece {full_lower.count('panel comercial')} veces"],
    }
    checks["plan8_figura8"] = {
        "desc": "Plan 8: epígrafe Figura 8 recorte declarado",
        "ok": "recorte" in full_lower or "diez registros" in full_lower or "10 registros" in full_lower,
        "evidence": find_context(paragraphs, "Figura 8")[:2] or find_context(paragraphs, "recorte")[:1],
    }
    checks["plan9_oraciones"] = {
        "desc": "Plan 9: reducción oraciones largas (cualitativo)",
        "ok": None,
        "evidence": ["Verificación cualitativa"],
    }
    checks["plan10_planilla"] = {
        "desc": "Plan 10 (opcional): planilla ejecuciones individuales",
        "ok": "planilla" in full_lower and ("anexo" in full_lower or "ausente" in full_lower or "limitación" in full_lower),
        "evidence": find_context(paragraphs, "planilla")[:2],
    }

    # === Q1-Q9 ===
    q1_pos = "condición de contraste" in full_lower and (
        "clasificador por reglas" in full_lower or "§5.4" in full_text or "5.4" in full_text
    )
    q1_neg = "no hay grupo de control" in full_lower and "condición de contraste" not in full_lower
    checks["Q1_anclaje_metodologico"] = {
        "desc": "Q1 OBLIG: §3.1 reformulado + condición de contraste",
        "ok": q1_pos and not (q1_neg and not q1_pos),
        "evidence": find_context(paragraphs, "condición de contraste")[:2] or find_context(paragraphs, "grupo de control")[:2],
    }

    s37 = section_paragraphs(paragraphs, [r"3\.7", r"procedimiento"])
    checks["Q1_3.7_item"] = {
        "desc": "Q1: §3.7 menciona campaña comparativa/clasificador reglas",
        "ok": bool(re.search(r"clasificador.*reglas|comparaci[oó]n.*pareada|simple-01|treinta conversaciones|30 conversaciones", s37, re.I)),
        "evidence": [s37[:500]] if s37 else [],
    }

    checks["Q2_ejemplo_ver"] = {
        "desc": "Q2 OBLIG: ejemplo §5.4 sin contradicción 'ver'",
        "ok": not bool(re.search(r"quiero ver uno.*no activan ninguna|no activan ninguna.*quiero ver", full_lower, re.DOTALL)),
        "evidence": find_context(paragraphs, "quiero ver")[:2] or find_context(paragraphs, "decisime que hay")[:2],
    }

    checks["Q3_TIB13_14"] = {
        "desc": "Q3 MUY REC: TIB-13/TIB-14 + fragilidad p-valor",
        "ok": ("tib-13" in full_lower or "tib-14" in full_lower) and (
            "frágil" in full_lower or "fragil" in full_lower or "discordantes" in full_lower
        ),
        "evidence": find_context(paragraphs, "TIB-13")[:2] or find_context(paragraphs, "discordantes")[:2],
    }

    checks["Q4_seis_comparaciones"] = {
        "desc": "Q4 MUY REC: 6 comparaciones §5.6 coherente o retiradas",
        "ok": bool(re.search(r"seis.*(ejecuciones|comparaciones|conversaciones)|no instrumentada|retira.*83", full_lower)),
        "evidence": find_context(paragraphs, "seis ejecuciones")[:2] or find_context(paragraphs, "no instrumentada")[:2],
    }

    checks["Q5_figura1_topologia"] = {
        "desc": "Q5 REC: topología Figura 1 (notificación bajo IF lead)",
        "ok": None,
        "evidence": find_context(paragraphs, "Figura 1")[:2],
    }

    kw_lists = [p["text"] for p in paragraphs if "palabras clave" in p["text"].lower()]
    checks["Q6_keywords_unificadas"] = {
        "desc": "Q6 REC: palabras clave portada y post-Resumen unificadas",
        "ok": full_lower.count("palabras clave") <= 2 and "panel comercial" in full_lower,
        "evidence": kw_lists[:4],
    }

    refs_section = ""
    in_refs = False
    for p in paragraphs:
        if re.match(r"^Referencias", p["text"], re.I):
            in_refs = True
        if in_refs:
            refs_section += p["text"] + "\n"
    mcnemar_idx = refs_section.lower().find("mcnemar")
    meta_idx = refs_section.lower().find("meta platforms")
    checks["Q7_mcnemar_orden"] = {
        "desc": "Q7 REC: McNemar antes de Meta Platforms",
        "ok": mcnemar_idx >= 0 and meta_idx >= 0 and mcnemar_idx < meta_idx,
        "evidence": [
            refs_section[mcnemar_idx : mcnemar_idx + 80] if mcnemar_idx >= 0 else "McNemar no encontrado",
            refs_section[meta_idx : meta_idx + 80] if meta_idx >= 0 else "Meta no encontrado",
        ],
    }

    checks["Q8_autorreferencia"] = {
        "desc": "Q8 REC: corregir autorreferencia §5.4",
        "ok": "sección 5.4" not in full_lower
        or "documentados más arriba" in full_lower
        or "más arriba en esta misma sección" in full_lower,
        "evidence": find_context(paragraphs, "sección 5.4")[:2] or find_context(paragraphs, "revisiones anteriores")[:2],
    }

    checks["Q9_tabla9"] = {
        "desc": "Q9 MENOR: Tabla 9 comparación §5.4",
        "ok": "tabla 9" in full_lower,
        "evidence": find_context(paragraphs, "Tabla 9")[:2],
    }

    checks["extra_5.4_mcnemar"] = {
        "desc": "Extra: §5.4 McNemar 66,7% vs 93,3%",
        "ok": "mcnemar" in full_lower
        and ("66,7" in full_text or "66.7" in full_text)
        and ("93,3" in full_text or "93.3" in full_text),
        "evidence": find_context(paragraphs, "McNemar")[:2],
    }

    zhao_idx = refs_section.lower().find("zhao")
    zheng_idx = refs_section.lower().find("zheng")
    checks["plan4_zhao_zheng"] = {
        "desc": "Plan 4: Zhao antes de Zheng",
        "ok": zhao_idx >= 0 and zheng_idx >= 0 and zhao_idx < zheng_idx,
        "evidence": [],
    }

    # v126 pendientes: mitiga parcialmente
    checks["v126_mitiga_parcialmente"] = {
        "desc": "Pendiente v126: 'mitiga parcialmente' corregido",
        "ok": "mitiga parcialmente" not in full_lower,
        "evidence": find_context(paragraphs, "mitiga parcialmente")[:2]
        or find_context(paragraphs, "no quedó instrumentada")[:2]
        or find_context(paragraphs, "Se había previsto mitigar")[:2],
    }

    checks["v126_38_62_coherencia"] = {
        "desc": "Pendiente v126: §3.8/§6.2 alineados con §5.6 (no instrumentada)",
        "ok": bool(
            re.search(
                r"no quedó instrumentada|no instrumentada|Se había previsto mitigar",
                section_paragraphs(paragraphs, [r"3\.8", r"validez"]) + section_paragraphs(paragraphs, [r"6\.2", r"generalizaciones"]),
                re.I,
            )
        ),
        "evidence": find_context(paragraphs, "no quedó instrumentada")[:2]
        + find_context(paragraphs, "Se había previsto")[:2],
    }

    return checks, refs_section


def main():
    pdf_text = extract_pdf_text(PDF_PATH)
    v127 = extract_docx_text(V127_PATH)
    v126 = extract_docx_text(V126_PATH) if V126_PATH.exists() else ""
    paras127 = extract_docx_paragraphs(V127_PATH)
    paras126 = extract_docx_paragraphs(V126_PATH) if V126_PATH.exists() else []

    (OUT_DIR / "devolucion4.txt").write_text(pdf_text, encoding="utf-8")
    (OUT_DIR / "v127.txt").write_text(v127, encoding="utf-8")
    if v126:
        (OUT_DIR / "v126.txt").write_text(v126, encoding="utf-8")

    checks, refs = run_devolucion4_checks(paras127, v127)

    # §5.6 analysis
    s56 = section_paragraphs(paras127, [r"5\.6", r"concordancia inter"])
    s56_v126 = section_paragraphs(paras126, [r"5\.6", r"concordancia inter"]) if paras126 else ""

    pendientes = {
        "5.6_text": s56[:2500],
        "has_833_claim": bool(re.search(r"83[,.]3\s*%?(?!\s*y\s*su\s*intervalo.*retira)", s56)) and "retira" not in s56.lower()[:200],
        "has_5_6_active": bool(re.search(r"5\s*/\s*6|5 de 6", s56, re.I)) and "retira" not in s56.lower(),
        "has_quiero_ver": "quiero ver" in s56.lower(),
        "mentions_n6_active": bool(re.search(r"n\s*=\s*6(?![^\n]{0,200}retira)", s56, re.I)),
        "declares_not_instrumented": bool(re.search(r"no instrumentada|no quedó instrumentada|retira.*83", s56, re.I)),
        "mitiga_parcialmente_count": v127.lower().count("mitiga parcialmente"),
        "no_quedo_instrumentada_count": len(find_all(r"no quedó instrumentada|no instrumentada", v127, re.I)),
    }

    # Global searches
    searches = {}
    for term in [
        "mitiga parcialmente",
        "Se había previsto mitigar",
        "no quedó instrumentada",
        "quiero ver uno",
        "decisime que hay",
        "n=6",
        "n = 6",
        "83,3",
        "5 de 6",
        "Figuras 5 a 8",
        "Figuras 3 a 6",
        "condición de contraste",
    ]:
        searches[term] = {
            "v127": v127.lower().count(term.lower()),
            "v126": v126.lower().count(term.lower()) if v126 else None,
        }

    # Figuras 5-8
    for fig in range(5, 9):
        pendientes[f"fig{fig}_ref"] = bool(re.search(rf"figura\s*{fig}\b", v127, re.I))

    # Diff v126 -> v127
    diff_hunks = []
    changed_lines = []
    if v126:
        v126_lines = v126.splitlines()
        v127_lines = v127.splitlines()
        diff_hunks = list(unified_diff(v126_lines, v127_lines, lineterm="", n=1))
        for line in diff_hunks:
            if line.startswith("+") and not line.startswith("+++"):
                changed_lines.append(line[1:])
            elif line.startswith("-") and not line.startswith("---"):
                changed_lines.append(f"[-] {line[1:]}")

    # Key phrase diff
    key_phrases = [
        "mitiga parcialmente",
        "Se había previsto mitigar",
        "no quedó instrumentada",
        "no instrumentada",
        "83,3",
        "n=6",
        "quiero ver",
        "decisime que hay",
        "condición de contraste",
        "clasificador por reglas",
    ]
    phrase_diff = {}
    for ph in key_phrases:
        phrase_diff[ph] = {
            "v126": len(re.findall(re.escape(ph), v126, re.I)) if v126 else None,
            "v127": len(re.findall(re.escape(ph), v127, re.I)),
        }

    # Score
    scored = [(k, c) for k, c in checks.items() if c["ok"] is not None]
    ok_count = sum(1 for _, c in scored if c["ok"])
    total = len(scored)
    pct = round(100 * ok_count / total, 1) if total else 0

    blockers = [c["desc"] for _, c in scored if not c["ok"] and "OBLIG" in c["desc"]]
    if pendientes["mitiga_parcialmente_count"] > 0:
        blockers.append("'mitiga parcialmente' aún presente")
    if pendientes["has_quiero_ver"]:
        blockers.append("§5.6 contiene 'quiero ver'")

    report = {
        "version": "v127",
        "checks": {k: {"desc": c["desc"], "ok": c["ok"], "evidence": c.get("evidence", [])} for k, c in checks.items()},
        "ok_count": ok_count,
        "total_scored": total,
        "pct": pct,
        "pendientes": pendientes,
        "searches": searches,
        "phrase_diff": phrase_diff,
        "blockers": blockers,
        "diff_line_count": len(diff_hunks),
        "diff_changed_sample": changed_lines[:40],
    }

    (OUT_DIR / "audit_report_v127.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    print("=" * 80)
    print("AUDITORÍA v127 vs Devolución 4")
    print("=" * 80)
    for k, c in checks.items():
        if c["ok"] is True:
            icon = "OK"
        elif c["ok"] is False:
            icon = "FAIL"
        else:
            icon = "WARN"
        print(f"[{icon}] {c['desc']}")
        for ev in (c.get("evidence") or [])[:1]:
            print(f"    -> {str(ev)[:180]}")

    print(f"\nCUMPLIMIENTO: {ok_count}/{total} = {pct}%")
    print(f"BLOCKERS: {blockers}")

    print("\n=== PENDIENTES v126 ===")
    print(json.dumps(pendientes, ensure_ascii=False, indent=2)[:4000])

    print("\n=== BÚSQUEDAS CLAVE ===")
    for term, counts in searches.items():
        if counts["v127"] or (counts["v126"] and counts["v126"] != counts["v127"]):
            print(f"  {term}: v126={counts['v126']} v127={counts['v127']}")

    print("\n=== PHRASE DIFF v126->v127 ===")
    for k, v in phrase_diff.items():
        if v["v126"] != v["v127"]:
            print(f"  {k}: {v['v126']} -> {v['v127']}")

    print("\n=== DIFF SAMPLE (first 30 changed) ===")
    for line in changed_lines[:30]:
        print(line[:200])

    print("\n=== §5.6 v127 (extract) ===")
    print(s56[:2000])
    if s56_v126 and s56 != s56_v126:
        print("\n=== §5.6 CHANGED vs v126 ===")
        s56_diff = list(unified_diff(s56_v126.splitlines(), s56.splitlines(), lineterm="", n=0))[:20]
        for l in s56_diff:
            print(l[:200])


if __name__ == "__main__":
    main()
